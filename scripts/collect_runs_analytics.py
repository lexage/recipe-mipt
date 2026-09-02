#!/usr/bin/env python3
"""Collect SWE-rebench experiment analytics from directories under ``runs``."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml


COLUMNS = (
    "run_name",
    "model_name",
    "total_instances",
    "resolved_instances",
    "unresolved_instances",
    "empty_patch_instances",
    "resolved_rate",
    "resolved_ids",
    "task_step_statistics",
    "critic_tool_name",
    "trajectory_summary_enabled",
    "critic_tool_call_count",
    "critic_tool_task_ids",
    "inference_duration_seconds",
    "evaluation_duration_seconds",
)

MAX_COLUMN_WIDTHS = {
    "run_name": 55,
    "model_name": 36,
    "total_instances": 18,
    "resolved_instances": 21,
    "unresolved_instances": 23,
    "empty_patch_instances": 25,
    "resolved_rate": 18,
    "resolved_ids": 50,
    "task_step_statistics": 70,
    "critic_tool_name": 22,
    "trajectory_summary_enabled": 30,
    "critic_tool_call_count": 25,
    "critic_tool_task_ids": 50,
    "inference_duration_seconds": 28,
    "evaluation_duration_seconds": 29,
}

CRITIC_TOOL_TYPES = {"CRITIC_TOOL", "SWE_REBENCH_CRITIC_TOOL"}
REPORT_COUNT_FIELDS = {
    "total_instances",
    "resolved_instances",
    "unresolved_instances",
    "empty_patch_instances",
}
INSTANCE_ID_PATTERN = re.compile(r"\binstance_id=([^\s\t]+)")
STEP_PATTERN = re.compile(r"\bSTEP\s+(\d+):", flags=re.IGNORECASE)
STAGE_PATTERN = re.compile(r"\bSTAGE[\t ]+([A-Z][A-Z0-9_]+)\b")
ERROR_TYPE_PATTERN = re.compile(
    r"^([A-Za-z_][A-Za-z0-9_.]*(?:Error|Exception)):\s*(.*)$"
)

IGNORED_LAST_STAGES = {
    "CONFIG_LOADED",
    "DATASET_LOADED",
    "DOCKER_CONTAINER_REMOVED",
    "PIPELINE_CLOSED",
    "RUN_FINISHED",
    "RUN_START",
    "TASKS_SELECTED",
    "WORKER_FINISHED",
}
FAILURE_STAGES = {"TASK_FAILED", "TASK_TIMEOUT", "WORKER_FAILED"}
EARLY_STAGES = {
    "TASK_SUBMITTED",
    "WORKER_STARTED",
    "IMAGE_RESOLVED",
    "WORKER_IMAGE_READY",
    "DOCKER_IMAGE_READY",
    "PIPELINE_CREATED",
    "DOCKER_CONTAINER_CREATED",
    "DOCKER_CHECKOUT_VALIDATED",
    "CONTAINER_STARTED",
    "PROMPT_BUILT",
    "PIPELINE_STARTED",
}

STATUS_RESOLVED = "решена"
STATUS_UNRESOLVED = "нерешена"
STATUS_EMPTY_PATCH = "пустой патч"
REPORT_TASK_ID_FIELDS = (
    "completed_ids",
    "incomplete_ids",
    "resolved_ids",
    "unresolved_ids",
    "empty_patch_ids",
    "error_ids",
)


@dataclass(frozen=True)
class PipelineAnalytics:
    model_name: str | None
    critic_tool_names: tuple[str, ...]
    trajectory_summary_enabled: bool


@dataclass(frozen=True)
class LogAnalytics:
    critic_call_count: int
    critic_task_ids: tuple[str, ...]
    task_steps: dict[str, int]
    task_diagnostics: dict[str, "TaskLogDiagnostics"]


@dataclass
class TaskLogDiagnostics:
    has_log: bool = False
    last_stage: str | None = None
    task_timeout: bool = False
    worker_failed: bool = False
    worker_error_type: str | None = None
    worker_error_message: str | None = None
    llm_request_error: str | None = None
    llm_response_error: str | None = None
    agent_termination_reason: str | None = None
    log_read_error: str | None = None


def warn(message: str) -> None:
    print(f"Предупреждение: {message}", file=sys.stderr)


def load_mapping(path: Path, *, yaml_file: bool = False) -> dict[str, Any] | None:
    try:
        with path.open(encoding="utf-8") as source:
            value = yaml.safe_load(source) if yaml_file else json.load(source)
    except (OSError, json.JSONDecodeError, yaml.YAMLError) as error:
        warn(f"не удалось прочитать {path}: {error}")
        return None

    if not isinstance(value, dict):
        warn(
            f"ожидался объект в {path}, "
            f"получено {type(value).__name__}"
        )
        return None
    return value


def as_mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def primary_agent_params(config: Mapping[str, Any]) -> Mapping[str, Any]:
    components = as_mapping(config.get("components"))
    for component_name in ("agent", "react_agent", "solver_agent"):
        component = as_mapping(components.get(component_name))
        params = as_mapping(component.get("params"))
        if params:
            return params
    return {}


def extract_pipeline_analytics(config: Mapping[str, Any]) -> PipelineAnalytics:
    components = as_mapping(config.get("components"))
    agent_params = primary_agent_params(config)

    model_name_value = agent_params.get("model_name")
    model_name = (
        str(model_name_value).strip()
        if model_name_value is not None and str(model_name_value).strip()
        else None
    )

    summary_value = agent_params.get("trajectory_summary_enabled", False)
    if isinstance(summary_value, str):
        trajectory_summary_enabled = summary_value.strip().lower() == "true"
    else:
        trajectory_summary_enabled = bool(summary_value)

    critic_tool_names: list[str] = []
    tools = components.get("tools")
    if isinstance(tools, list):
        for tool in tools:
            tool_config = as_mapping(tool)
            tool_type = str(tool_config.get("type", "")).strip().upper()
            if tool_type not in CRITIC_TOOL_TYPES:
                continue
            params = as_mapping(tool_config.get("params"))
            name_value = params.get("name", params.get("method", "critic"))
            name = str(name_value).strip()
            if name and name not in critic_tool_names:
                critic_tool_names.append(name)

    return PipelineAnalytics(
        model_name=model_name,
        critic_tool_names=tuple(critic_tool_names),
        trajectory_summary_enabled=trajectory_summary_enabled,
    )


def find_evaluation_report(run_dir: Path) -> dict[str, Any] | None:
    evaluation_dir = run_dir / "evaluation"
    if not evaluation_dir.is_dir():
        warn(
            f"в запуске {run_dir.name} отсутствует папка evaluation"
        )
        return None

    reports: list[tuple[Path, dict[str, Any]]] = []
    for path in sorted(evaluation_dir.glob("*.json")):
        if path.name == "evaluation_metadata.json":
            continue
        data = load_mapping(path)
        if data is not None and REPORT_COUNT_FIELDS.issubset(data):
            reports.append((path, data))

    if not reports:
        warn(
            f"для запуска {run_dir.name} не найден итоговый JSON оценки"
        )
        return None

    exact = [
        report
        for report in reports
        if report[0].name.endswith(f".{run_dir.name}.json")
    ]
    matching = exact or [report for report in reports if run_dir.name in report[0].name]
    selected = matching or reports
    if len(selected) > 1:
        warn(
            f"для запуска {run_dir.name} найдено несколько отчётов "
            "оценки; "
            f"используется {selected[0][0].name}"
        )
    return selected[0][1]


def infer_instance_id_from_filename(path: Path) -> str | None:
    marker = "swe-rebench-"
    if marker not in path.stem:
        return None
    instance_id = path.stem.rsplit(marker, maxsplit=1)[1].strip()
    return instance_id or None


def build_action_pattern(critic_tool_names: Iterable[str]) -> re.Pattern[str] | None:
    names = sorted(
        {name.strip() for name in critic_tool_names if name.strip()},
        key=len,
        reverse=True,
    )
    if not names:
        return None
    alternatives = "|".join(re.escape(name) for name in names)
    return re.compile(
        rf"(?<![A-Z0-9_])ACTION:\s*(?:{alternatives})(?![A-Z0-9_.-])",
        flags=re.IGNORECASE,
    )


def extract_log_field(line: str, field_name: str) -> str | None:
    match = re.search(
        rf"(?:^|\t){re.escape(field_name)}=([^\t\r\n]*)",
        line,
    )
    if not match:
        return None
    value = match.group(1).strip()
    return value or None


def get_task_diagnostics(
    diagnostics: dict[str, TaskLogDiagnostics], instance_id: str
) -> TaskLogDiagnostics:
    task_diagnostics = diagnostics.setdefault(
        instance_id, TaskLogDiagnostics()
    )
    task_diagnostics.has_log = True
    return task_diagnostics


def collect_log_analytics(
    logs_dir: Path, critic_tool_names: Iterable[str]
) -> LogAnalytics:
    action_pattern = build_action_pattern(critic_tool_names)
    if not logs_dir.is_dir():
        warn(f"не найдена папка логов {logs_dir}")
        return LogAnalytics(0, (), {}, {})

    call_count = 0
    task_ids: set[str] = set()
    task_steps: dict[str, int] = {}
    task_diagnostics: dict[str, TaskLogDiagnostics] = {}
    unattributed_calls = 0
    unattributed_steps = 0

    for path in sorted(logs_dir.rglob("*.log")):
        file_instance_id = infer_instance_id_from_filename(path)
        current_instance_id = file_instance_id
        if file_instance_id:
            get_task_diagnostics(task_diagnostics, file_instance_id)
        try:
            source = path.open(encoding="utf-8", errors="replace")
        except OSError as error:
            warn(f"не удалось прочитать лог {path}: {error}")
            if file_instance_id:
                diagnostics = get_task_diagnostics(
                    task_diagnostics, file_instance_id
                )
                diagnostics.log_read_error = f"{type(error).__name__}: {error}"
            continue

        with source:
            for line in source:
                instance_match = INSTANCE_ID_PATTERN.search(line)
                if instance_match:
                    current_instance_id = instance_match.group(1)
                    get_task_diagnostics(
                        task_diagnostics, current_instance_id
                    )

                diagnostics = (
                    get_task_diagnostics(
                        task_diagnostics, current_instance_id
                    )
                    if current_instance_id
                    else None
                )

                stage_match = STAGE_PATTERN.search(line)
                if diagnostics is not None and stage_match:
                    stage = stage_match.group(1)
                    if (
                        stage not in IGNORED_LAST_STAGES
                        and (
                            diagnostics.last_stage not in FAILURE_STAGES
                            or stage in FAILURE_STAGES
                        )
                    ):
                        diagnostics.last_stage = stage
                    if stage == "TASK_TIMEOUT":
                        diagnostics.task_timeout = True
                    elif stage in {"WORKER_FAILED", "TASK_FAILED"}:
                        diagnostics.worker_failed = True
                        error_type = extract_log_field(line, "error_type")
                        message = extract_log_field(line, "message")
                        if error_type:
                            diagnostics.worker_error_type = error_type
                        if message:
                            diagnostics.worker_error_message = message
                    elif stage == "AGENT_TERMINATED":
                        reason = extract_log_field(line, "reason")
                        if reason:
                            diagnostics.agent_termination_reason = reason

                if diagnostics is not None:
                    if "LLM_REQUEST_FAILED" in line:
                        error = extract_log_field(line, "error")
                        if error:
                            diagnostics.llm_request_error = error
                    elif "LLM_RESPONSE_INVALID" in line:
                        error = extract_log_field(line, "error")
                        if error:
                            diagnostics.llm_response_error = error

                step_matches = list(STEP_PATTERN.finditer(line))
                if step_matches:
                    if current_instance_id:
                        task_steps[current_instance_id] = int(
                            step_matches[-1].group(1)
                        )
                    else:
                        unattributed_steps += len(step_matches)

                if action_pattern is not None:
                    action_matches = list(action_pattern.finditer(line))
                    call_count += len(action_matches)
                    if action_matches:
                        if current_instance_id:
                            task_ids.add(current_instance_id)
                        else:
                            unattributed_calls += len(action_matches)

    if unattributed_calls:
        warn(
            f"для {unattributed_calls} вызовов критики в {logs_dir} "
            "не удалось определить instance_id"
        )
    if unattributed_steps:
        warn(
            f"для {unattributed_steps} упоминаний STEP в {logs_dir} "
            "не удалось определить instance_id"
        )
    return LogAnalytics(
        call_count,
        tuple(sorted(task_ids)),
        task_steps,
        task_diagnostics,
    )


def normalize_ids(value: Any, field_name: str) -> list[str]:
    if value is None:
        return []
    if not isinstance(value, (list, tuple, set)):
        warn(
            f"в поле {field_name} ожидался список идентификаторов, "
            f"получено {type(value).__name__}"
        )
        value = [value]

    ids: list[str] = []
    for item in value:
        instance_id = str(item).strip()
        if instance_id:
            ids.append(instance_id)
    return sorted(ids) if isinstance(value, set) else ids


def unique_instance_ids(values: Iterable[str], source: Path) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    duplicate_count = 0
    for value in values:
        instance_id = str(value).strip()
        if not instance_id:
            continue
        if instance_id in seen:
            duplicate_count += 1
            continue
        seen.add(instance_id)
        result.append(instance_id)

    if duplicate_count:
        warn(
            f"в {source} пропущено повторяющихся instance_id: "
            f"{duplicate_count}"
        )
    return result


def load_instance_ids_file(path: Path) -> list[str]:
    try:
        with path.open(encoding="utf-8") as source:
            ids = unique_instance_ids(source, path)
    except OSError as error:
        warn(f"не удалось прочитать список задач {path}: {error}")
        return []

    if not ids:
        warn(f"список задач пуст: {path}")
    return ids


def load_task_jsonl_instance_ids(path: Path) -> list[str]:
    instance_ids: list[str] = []
    try:
        source = path.open(encoding="utf-8")
    except OSError as error:
        warn(f"не удалось прочитать датасет запуска {path}: {error}")
        return []

    with source:
        for line_number, line in enumerate(source, start=1):
            if not line.strip():
                continue
            try:
                task = json.loads(line)
            except json.JSONDecodeError as error:
                warn(f"невалидный JSON в {path}:{line_number}: {error}")
                continue
            if not isinstance(task, Mapping):
                warn(f"ожидался объект в {path}:{line_number}")
                continue
            instance_id = task.get("instance_id")
            if instance_id is None or not str(instance_id).strip():
                warn(f"в {path}:{line_number} отсутствует instance_id")
                continue
            instance_ids.append(str(instance_id))

    ids = unique_instance_ids(instance_ids, path)
    if not ids:
        warn(f"в датасете запуска нет задач: {path}")
    return ids


def collect_run_instance_ids(
    run_dir: Path,
    report: Mapping[str, Any],
    task_steps: Mapping[str, int],
    error_instance_ids: Iterable[str],
) -> list[str]:
    candidates: list[tuple[Path, list[str]]] = []
    instance_ids_path = run_dir / "instance_ids.txt"
    if instance_ids_path.is_file():
        ids = load_instance_ids_file(instance_ids_path)
        if ids:
            candidates.append((instance_ids_path, ids))

    task_jsonl_path = run_dir / "task.jsonl"
    if task_jsonl_path.is_file():
        ids = load_task_jsonl_instance_ids(task_jsonl_path)
        if ids:
            candidates.append((task_jsonl_path, ids))

    total_instances = report.get("total_instances")
    expected_count = (
        total_instances
        if isinstance(total_instances, int)
        and not isinstance(total_instances, bool)
        and total_instances > 0
        else None
    )
    if candidates:
        if expected_count is not None:
            for _, ids in candidates:
                if len(ids) == expected_count:
                    return ids
        selected_path, selected_ids = max(
            candidates, key=lambda candidate: len(candidate[1])
        )
        if expected_count is not None:
            warn(
                f"в {selected_path} найдено задач: {len(selected_ids)}, "
                f"а в отчёте указано total_instances={expected_count}"
            )
        return selected_ids

    fallback_ids = set(task_steps)
    fallback_ids.update(error_instance_ids)
    for field in REPORT_TASK_ID_FIELDS:
        fallback_ids.update(normalize_ids(report.get(field), field))
    selected_ids = sorted(fallback_ids)
    if expected_count is not None and len(selected_ids) != expected_count:
        warn(
            f"для запуска {run_dir.name} удалось определить "
            f"{len(selected_ids)} из {expected_count} instance_id; "
            "файлы instance_ids.txt и task.jsonl отсутствуют или пусты"
        )
    return selected_ids


def load_error_records(path: Path) -> dict[str, dict[str, Any]]:
    if not path.is_file():
        return {}

    records: dict[str, dict[str, Any]] = {}
    try:
        source = path.open(encoding="utf-8")
    except OSError as error:
        warn(f"не удалось прочитать {path}: {error}")
        return {}

    with source:
        for line_number, line in enumerate(source, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as error:
                warn(f"невалидный JSON в {path}:{line_number}: {error}")
                continue
            if not isinstance(record, dict):
                warn(f"ожидался объект в {path}:{line_number}")
                continue
            instance_id_value = record.get("instance_id")
            instance_id = (
                str(instance_id_value).strip()
                if instance_id_value is not None
                else ""
            )
            if not instance_id:
                warn(f"в {path}:{line_number} отсутствует instance_id")
                continue
            records[instance_id] = record
    return records


def compact_text(value: Any, limit: int = 300) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split())
    if not text:
        return None
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def split_error_text(value: Any) -> tuple[str | None, str | None]:
    text = compact_text(value)
    if text is None:
        return None, None
    match = ERROR_TYPE_PATTERN.match(text)
    if not match:
        return None, text
    error_type = match.group(1)
    message = compact_text(match.group(2)) or text
    return error_type, message


def build_steps_null_reason(
    diagnostics: TaskLogDiagnostics | None,
    error_record: Mapping[str, Any],
) -> dict[str, Any]:
    details = as_mapping(error_record.get("details"))
    runner_error_type = compact_text(error_record.get("error_type"))
    runner_message = compact_text(error_record.get("message"))
    termination_reason = compact_text(details.get("termination_reason"))
    agent_termination_reason = compact_text(
        diagnostics.agent_termination_reason if diagnostics else None
    ) or compact_text(details.get("agent_termination_reason"))
    last_stage = diagnostics.last_stage if diagnostics else None

    error_type: str | None = None
    message: str | None = None
    if (
        (diagnostics is not None and diagnostics.task_timeout)
        or termination_reason == "task_timeout"
    ):
        code = "task_timeout"
        error_type = runner_error_type or "TimeoutError"
        message = runner_message or (
            "Задача завершилась по таймауту до первого записанного STEP."
        )
    elif diagnostics is not None and diagnostics.llm_request_error:
        code = "llm_request_failed"
        error_type, message = split_error_text(
            diagnostics.llm_request_error
        )
    elif diagnostics is not None and diagnostics.llm_response_error:
        code = "llm_response_invalid"
        error_type, message = split_error_text(
            diagnostics.llm_response_error
        )
    elif runner_error_type == "EmptyModelPatchError":
        code = "empty_model_patch"
        error_type = runner_error_type
        message = runner_message
    elif diagnostics is not None and diagnostics.worker_failed:
        code = "worker_failed"
        error_type = runner_error_type or diagnostics.worker_error_type
        message = runner_message or diagnostics.worker_error_message
    elif error_record:
        code = "inference_error"
        error_type = runner_error_type
        message = runner_message
    elif diagnostics is not None and diagnostics.log_read_error:
        code = "log_unreadable"
        error_type, message = split_error_text(diagnostics.log_read_error)
    elif agent_termination_reason:
        code = "agent_terminated_before_first_step"
        message = agent_termination_reason
    elif diagnostics is not None and last_stage in EARLY_STAGES:
        code = "stopped_before_first_step"
        message = (
            "Выполнение остановилось до первого записанного STEP; "
            f"последнее событие: {last_stage}."
        )
    elif diagnostics is not None and diagnostics.has_log:
        code = "no_step_marker"
        message = "В логе задачи нет ни одной строки вида STEP N:."
    else:
        code = "no_log"
        message = "Для задачи не найден лог."

    reason: dict[str, Any] = {"code": code}
    if last_stage:
        reason["last_stage"] = last_stage
    if error_type:
        reason["error_type"] = error_type
    if message:
        reason["message"] = message
    if runner_error_type and runner_error_type != error_type:
        reason["runner_error_type"] = runner_error_type
    if termination_reason:
        reason["termination_reason"] = termination_reason
    if agent_termination_reason:
        reason["agent_termination_reason"] = agent_termination_reason
    return reason


def serialize_ids(value: Any) -> str:
    if value is None:
        return ""
    return json.dumps(normalize_ids(value, "ids"), ensure_ascii=False)


def calculate_resolved_rate(report: Mapping[str, Any]) -> float | None:
    total_instances = report.get("total_instances")
    resolved_instances = report.get("resolved_instances")
    if isinstance(total_instances, bool) or isinstance(resolved_instances, bool):
        return None
    if not isinstance(total_instances, (int, float)):
        return None
    if not isinstance(resolved_instances, (int, float)):
        return None
    if total_instances <= 0:
        return None
    return resolved_instances / total_instances


def serialize_task_step_statistics(
    instance_ids: Iterable[str],
    report: Mapping[str, Any],
    task_steps: Mapping[str, int],
    task_diagnostics: Mapping[str, TaskLogDiagnostics],
    error_records: Mapping[str, Mapping[str, Any]],
) -> str:
    ids_by_field = {
        field: set(normalize_ids(report.get(field), field))
        for field in REPORT_TASK_ID_FIELDS
    }
    all_ids = set(instance_ids)

    status_by_id = {instance_id: STATUS_UNRESOLVED for instance_id in all_ids}
    for instance_id in ids_by_field["resolved_ids"]:
        status_by_id[instance_id] = STATUS_RESOLVED
    for instance_id in ids_by_field["empty_patch_ids"]:
        status_by_id[instance_id] = STATUS_EMPTY_PATCH

    statistics = {
        instance_id: {
            "status": status_by_id[instance_id],
            "steps": task_steps.get(instance_id),
            "steps_null_reason": (
                None
                if instance_id in task_steps
                else build_steps_null_reason(
                    task_diagnostics.get(instance_id),
                    error_records.get(instance_id, {}),
                )
            ),
        }
        for instance_id in sorted(all_ids)
    }
    return json.dumps(
        statistics,
        ensure_ascii=False,
        separators=(",", ":"),
    )


def nested_value(data: Mapping[str, Any], section: str, field: str) -> Any:
    return as_mapping(data.get(section)).get(field)


def collect_run(run_dir: Path) -> dict[str, Any]:
    pipeline_path = run_dir / "pipeline.yaml"
    config = (
        load_mapping(pipeline_path, yaml_file=True)
        if pipeline_path.is_file()
        else None
    )
    if config is None:
        warn(
            f"для запуска {run_dir.name} не удалось прочитать pipeline.yaml"
        )
        pipeline = PipelineAnalytics(None, (), False)
    else:
        pipeline = extract_pipeline_analytics(config)

    report = find_evaluation_report(run_dir) or {}

    timings_path = run_dir / "timings.json"
    timings = load_mapping(timings_path) if timings_path.is_file() else None
    if timings is None:
        warn(
            f"для запуска {run_dir.name} не удалось прочитать timings.json"
        )
        timings = {}

    log_analytics = collect_log_analytics(
        run_dir / "logs", pipeline.critic_tool_names
    )
    error_records = load_error_records(run_dir / "errors.jsonl")
    instance_ids = collect_run_instance_ids(
        run_dir,
        report,
        log_analytics.task_steps,
        error_records,
    )

    return {
        "run_name": run_dir.name,
        "model_name": pipeline.model_name,
        "total_instances": report.get("total_instances"),
        "resolved_instances": report.get("resolved_instances"),
        "unresolved_instances": report.get("unresolved_instances"),
        "empty_patch_instances": report.get("empty_patch_instances"),
        "resolved_rate": calculate_resolved_rate(report),
        "resolved_ids": serialize_ids(report.get("resolved_ids")),
        "task_step_statistics": serialize_task_step_statistics(
            instance_ids,
            report,
            log_analytics.task_steps,
            log_analytics.task_diagnostics,
            error_records,
        ),
        "critic_tool_name": ", ".join(pipeline.critic_tool_names),
        "trajectory_summary_enabled": pipeline.trajectory_summary_enabled,
        "critic_tool_call_count": log_analytics.critic_call_count,
        "critic_tool_task_ids": serialize_ids(log_analytics.critic_task_ids),
        "inference_duration_seconds": nested_value(
            timings, "inference", "duration_seconds"
        ),
        "evaluation_duration_seconds": nested_value(
            timings, "evaluation", "duration_seconds"
        ),
    }


def collect_runs(runs_dir: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for run_dir in sorted(path for path in runs_dir.iterdir() if path.is_dir()):
        if not (run_dir / "pipeline.yaml").is_file():
            warn(f"папка {run_dir.name} пропущена: нет pipeline.yaml")
            continue
        rows.append(collect_run(run_dir))
    return rows


def csv_value(value: Any) -> Any:
    if isinstance(value, bool):
        return str(value).lower()
    return "" if value is None else value


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow({column: csv_value(row.get(column)) for column in COLUMNS})


def write_excel(path: Path, rows: list[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError(
            "Для записи Excel установите зависимости проекта, "
            "включая openpyxl"
        ) from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "runs"
    sheet.append(list(COLUMNS))
    for row in rows:
        sheet.append([row.get(column) for column in COLUMNS])

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.row_dimensions[1].height = 32
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    numeric_columns = {
        "total_instances",
        "resolved_instances",
        "unresolved_instances",
        "empty_patch_instances",
        "critic_tool_call_count",
        "inference_duration_seconds",
        "evaluation_duration_seconds",
    }
    percentage_columns = {"resolved_rate"}
    for column_index, column_name in enumerate(COLUMNS, start=1):
        column_letter = get_column_letter(column_index)
        values = [column_name, *(str(row.get(column_name) or "") for row in rows)]
        width = min(
            max(len(value) for value in values) + 2,
            MAX_COLUMN_WIDTHS[column_name],
        )
        sheet.column_dimensions[column_letter].width = max(width, 12)
        if column_name in numeric_columns:
            for cell in sheet[column_letter][1:]:
                cell.number_format = "#,##0"
        elif column_name in percentage_columns:
            for cell in sheet[column_letter][1:]:
                cell.number_format = "0.0%"

    wrapped_columns = {
        "run_name",
        "model_name",
        "resolved_ids",
        "task_step_statistics",
        "trajectory_summary_enabled",
        "critic_tool_task_ids",
    }
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=COLUMNS[cell.column - 1] in wrapped_columns,
            )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Собрать сводную аналитику запусков SWE-rebench "
            "из папки runs "
            "в CSV и Excel."
        )
    )
    parser.add_argument(
        "--runs-dir",
        type=Path,
        default=Path("runs"),
        help=(
            "Папка с экспериментами (по умолчанию: runs)"
        ),
    )
    parser.add_argument(
        "--output-prefix",
        type=Path,
        default=None,
        help=(
            "Путь без расширения для выходных файлов; "
            "по умолчанию "
            "<runs-dir>/runs_analytics"
        ),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    runs_dir = args.runs_dir.expanduser().resolve()
    if not runs_dir.is_dir():
        print(
            f"Папка с запусками не найдена: {runs_dir}",
            file=sys.stderr,
        )
        return 2

    output_prefix = args.output_prefix or runs_dir / "runs_analytics"
    output_prefix = output_prefix.expanduser().resolve()
    if output_prefix.suffix.lower() in {".csv", ".xlsx"}:
        output_prefix = output_prefix.with_suffix("")

    rows = collect_runs(runs_dir)
    csv_path = output_prefix.with_suffix(".csv")
    excel_path = output_prefix.with_suffix(".xlsx")
    write_csv(csv_path, rows)
    try:
        write_excel(excel_path, rows)
    except RuntimeError as error:
        print(str(error), file=sys.stderr)
        return 1

    print(f"Обработано запусков: {len(rows)}")
    print(f"CSV: {csv_path}")
    print(f"Excel: {excel_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
